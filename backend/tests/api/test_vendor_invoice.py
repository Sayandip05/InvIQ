"""
Tests for Vendor Delivery Invoice Generation, ReportLab PDF Rendering, and Azure Blob Storage.
"""

import pytest
from datetime import date
from io import BytesIO
from unittest.mock import patch, MagicMock
import openpyxl

from app.infrastructure.database.models import Item, Location, User, VendorUpload, VendorInvoice
from app.infrastructure.database.invoice_repo import InvoiceRepository
from app.application.invoice_pdf_service import InvoicePdfService
from app.application.vendor_service import VendorService
from app.infrastructure.storage.azure_blob_storage import AzureBlobStorageService
from tests.conftest import get_auth_header


def _create_sample_excel(items_data: list[tuple[str, int, float]]) -> bytes:
    """Helper to create sample delivery Excel bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deliveries"

    ws.append(["item_name", "quantity_received", "unit_price", "delivery_date", "notes"])
    for name, qty, price in items_data:
        ws.append([name, qty, price, "2026-08-14", "Delivered on time"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Invoice Number Generation ─────────────────────────────────────────────

def test_generate_invoice_number(db, admin_user):
    repo = InvoiceRepository(db)
    test_date = date(2029, 11, 25)
    inv1 = repo.generate_next_invoice_number(test_date)
    assert inv1 == "INV-20291125-001"

    user = db.query(User).filter(User.username == admin_user["username"]).first()
    loc = db.query(Location).first()
    if not loc:
        loc = Location(org_id=1, name="Main Depot", type="warehouse", region="East")
        db.add(loc)
        db.commit()

    upload = VendorUpload(
        org_id=1,
        vendor_user_id=user.id,
        filename="test.xlsx",
        location_id=loc.id,
        total_rows=1,
        success_rows=1,
        status="COMPLETED",
    )
    db.add(upload)
    db.commit()

    repo.create(
        org_id=1,
        vendor_user_id=user.id,
        vendor_upload_id=upload.id,
        invoice_number=inv1,
        invoice_date=test_date,
        line_items=[{"item_name": "Paracetamol", "qty": 10, "unit_price": 50.0, "total": 500.0}],
        subtotal=500.0,
        tax_amount=90.0,
        total_amount=590.0,
    )

    inv2 = repo.generate_next_invoice_number(test_date)
    assert inv2 == "INV-20291125-002"

    # Test collision resilience: if repo.create is given an already existing invoice_number, it auto-retries and succeeds
    upload2 = VendorUpload(
        org_id=1,
        vendor_user_id=user.id,
        filename="test2.xlsx",
        location_id=loc.id,
        total_rows=1,
        success_rows=1,
        status="COMPLETED",
    )
    db.add(upload2)
    db.commit()

    # Pass inv1 intentionally (simulating simultaneous request collision)
    invoice_collided = repo.create(
        org_id=1,
        vendor_user_id=user.id,
        vendor_upload_id=upload2.id,
        invoice_number=inv1,  # collision
        invoice_date=test_date,
        line_items=[{"item_name": "Ibuprofen", "qty": 5, "unit_price": 20.0, "total": 100.0}],
        subtotal=100.0,
        tax_amount=18.0,
        total_amount=118.0,
    )
    assert invoice_collided.invoice_number == "INV-20291125-002"




# ── ReportLab PDF Generation ──────────────────────────────────────────────

def test_invoice_pdf_rendering():
    invoice_data = {
        "invoice_number": "INV-20260814-001",
        "invoice_date": "2026-08-14",
        "line_items": [
            {"item_name": "Paracetamol 500mg", "quantity": 100, "unit": "Tablets", "unit_price": 2.50, "total": 250.0},
            {"item_name": "Insulin Regular 100IU", "quantity": 20, "unit": "Vials", "unit_price": 350.0, "total": 7000.0},
        ],
        "subtotal": 7250.0,
        "tax_amount": 1305.0,
        "total_amount": 8555.0,
        "status": "ISSUED",
    }
    vendor_data = {
        "username": "medicorp_vendor",
        "full_name": "MediCorp India Ltd",
        "email": "deliveries@medicorp.in",
    }
    location_data = {
        "name": "Central Pharmacy Warehouse",
        "region": "Kolkata North",
    }

    pdf_bytes = InvoicePdfService.generate_invoice_pdf(
        invoice_data=invoice_data,
        vendor_data=vendor_data,
        location_data=location_data,
    )

    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    assert pdf_bytes[:5] == b"%PDF-"


# ── Azure Blob Storage Service ───────────────────────────────────────────

def test_azure_blob_storage_service_upload_and_download():
    with patch("azure.storage.blob.BlobServiceClient.from_connection_string") as mock_bsc:
        mock_client = MagicMock()
        mock_container = MagicMock()
        mock_blob = MagicMock()

        mock_bsc.return_value = mock_client
        mock_client.get_container_client.return_value = mock_container
        mock_container.get_blob_client.return_value = mock_blob
        mock_blob.url = "https://inviqstorage.blob.core.windows.net/inviq-documents/invoices/2026/08/INV-001.pdf"
        mock_blob.download_blob.return_value.readall.return_value = b"%PDF-mock-bytes"

        with patch("app.infrastructure.storage.azure_blob_storage.settings") as s:
            s.AZURE_STORAGE_CONNECTION_STRING = "DefaultEndpointsProtocol=https;AccountName=test;AccountKey=test;EndpointSuffix=core.windows.net"
            s.AZURE_STORAGE_CONTAINER_NAME = "inviq-documents"
            s.AZURE_STORAGE_ACCOUNT_NAME = None
            s.AZURE_STORAGE_ACCOUNT_KEY = None

            storage = AzureBlobStorageService()
            assert storage.is_available is True

            url = storage.upload_file(b"%PDF-test", "invoices/2026/08/INV-001.pdf")
            assert url == mock_blob.url

            downloaded = storage.download_file("invoices/2026/08/INV-001.pdf")
            assert downloaded == b"%PDF-mock-bytes"


# ── Vendor Upload with Auto Invoice Generation ───────────────────────────

def test_vendor_upload_creates_invoice(db):
    from app.core.security import hash_password

    loc = db.query(Location).first()
    if not loc:
        loc = Location(org_id=1, name="Central Depot", type="warehouse", region="Kolkata")
        db.add(loc)
        db.commit()

    item = db.query(Item).first()
    if not item:
        item = Item(org_id=1, name="Amoxicillin 500mg", category="Antibiotics", unit="Capsules", lead_time_days=3, min_stock=50)
        db.add(item)
        db.commit()

    vendor = db.query(User).filter(User.role == "vendor").first()
    if not vendor:
        vendor = User(
            username="test_vendor",
            email="vendor@test.com",
            hashed_password=hash_password("VendorPass123!"),
            role="vendor",
            is_active=True,
        )
        db.add(vendor)
        db.commit()

    excel_bytes = _create_sample_excel([(item.name, 46, 120.0)])

    service = VendorService(db)
    res = service.parse_and_process_excel(
        file_content=excel_bytes,
        filename="vendor_delivery_46items.xlsx",
        location_id=loc.id,
        vendor_user_id=vendor.id,
    )

    assert res["success"] is True
    data = res["data"]
    assert data["success"] == 1
    assert "invoice" in data

    invoice_info = data["invoice"]
    assert invoice_info["items_count"] == 1
    assert invoice_info["subtotal"] == 5520.0  # 46 * 120
    assert invoice_info["tax_amount"] == 993.60  # 5520 * 0.18
    assert invoice_info["total_amount"] == 6513.60
    assert invoice_info["status"] == "ISSUED"

    # Verify in DB
    inv = db.query(VendorInvoice).filter(VendorInvoice.id == invoice_info["invoice_id"]).first()
    assert inv is not None
    assert inv.vendor_upload_id == data["upload_id"]
    assert inv.pdf_content is not None
    assert inv.pdf_content[:5] == b"%PDF-"


# ── API Endpoints ─────────────────────────────────────────────────────────

def test_vendor_invoice_api_endpoints(client, db, admin_user):
    headers = get_auth_header(client, admin_user["username"], admin_user["password"])

    # 1. List invoices
    res_list = client.get("/api/vendor/invoices", headers=headers)
    assert res_list.status_code == 200
    list_data = res_list.json()
    assert list_data["success"] is True
    assert "data" in list_data

    if list_data["total"] > 0:
        inv_id = list_data["data"][0]["id"]

        # 2. Get detail
        res_detail = client.get(f"/api/vendor/invoices/{inv_id}", headers=headers)
        assert res_detail.status_code == 200
        detail_data = res_detail.json()["data"]
        assert "line_items" in detail_data
        assert "total_amount" in detail_data

        # 3. Download PDF
        res_pdf = client.get(f"/api/vendor/invoices/{inv_id}/pdf", headers=headers)
        assert res_pdf.status_code == 200
        assert res_pdf.headers["content-type"] == "application/pdf"
        assert res_pdf.content[:5] == b"%PDF-"


def test_download_standard_template_endpoint(client, admin_user):
    """Test downloading the official standardized 11-column Excel delivery template."""
    headers = get_auth_header(client, admin_user["username"], admin_user["password"])
    res = client.get("/api/vendor/template", headers=headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "InvIQ_Medicine_Delivery_Template.xlsx" in res.headers["content-disposition"]

    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(res.content))
    ws = wb.active
    assert ws.title == "Medicine Delivery Manifest"

    # Verify all 11 standard headers
    expected_headers = [
        "item_name", "quantity", "unit", "batch_number", "expiry_date",
        "purchase_rate", "mrp", "category", "storage_temp", "delivery_date", "invoice_no"
    ]
    actual_headers = [cell.value for cell in ws[1]]
    assert actual_headers == expected_headers

    # Verify sample row has real medicine data
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == "Paracetamol 500mg"
    assert row2[1] == 100
    assert row2[2] == "strip"
    assert row2[3] == "BT-2026-A1"


